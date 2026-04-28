"""把 4181 的审批件数同步进 allocation_*.xlsx (Bug B'' 一致性).

action:
- 扫 SCRIPT_DIR/allocation_*.xlsx
- 每个文件的 B1=offer_id, R17 headers 含 "集群" + "(合计件)"
- 从 4181 拉对应 sku 最新 approved/dispatched plan, 取 allocations[].pieces
- 在原 R17 末尾追加 "审批件数" 列, 写到 R18+ 各 cluster 行
- 同时追加 "审批 vs solver" 列 = 审批件 - 合计件 (+/-)
- 不改原文件, 写到 allocation_<sku>_<date>_with_approved.xlsx
- 同时 print stdout 概览

用法: python3 sync_allocation_approved.py [--account 丝绸生活]
"""
from __future__ import annotations

import argparse
import json
import urllib.parse
import urllib.request
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

SCRIPT_DIR = Path(__file__).resolve().parent
RS = "http://127.0.0.1:4181"


def fetch_approved_plan_for_sku(account: str, offer_id: str) -> dict[str, int]:
    enc = urllib.parse.quote(account, safe="")
    with urllib.request.urlopen(f"{RS}/plans?account={enc}&limit=200", timeout=30) as f:
        lst = json.load(f)
    chosen = None
    for meta in lst:
        if meta.get("sku") != offer_id:
            continue
        if meta.get("status") not in ("approved", "dispatched"):
            continue
        if chosen is None or (meta.get("created_at", "") > chosen.get("created_at", "")):
            chosen = meta
    if not chosen:
        return {}
    with urllib.request.urlopen(f"{RS}/plans/{chosen['plan_id']}", timeout=30) as f:
        p = json.load(f)
    out: dict[str, int] = {}
    for a in (p.get("allocations") or []):
        cl = a.get("cluster")
        if cl:
            out[cl] = int(a.get("pieces") or 0)
    return out


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--account", default="丝绸生活")
    args = ap.parse_args()

    files = sorted(SCRIPT_DIR.glob("allocation_*.xlsx"))
    files = [f for f in files if "_with_approved" not in f.name]
    if not files:
        print("未找到 allocation_*.xlsx")
        return 1

    fill_yellow = PatternFill("solid", fgColor="FFE9C2")
    fill_red = PatternFill("solid", fgColor="F0D6D6")
    fill_green = PatternFill("solid", fgColor="D6F0C8")
    bold = Font(bold=True)

    summaries = []
    for src in files:
        wb = load_workbook(src)
        ws = wb.active
        offer_id = ws["B1"].value
        if not offer_id or not isinstance(offer_id, str):
            print(f"跳 {src.name}: B1 不是 offer_id")
            continue

        approved = fetch_approved_plan_for_sku(args.account, offer_id)
        if not approved:
            print(f"跳 {src.name}: 4181 无 {offer_id} 审批 plan")
            continue

        hdr = [c.value for c in ws[17]]
        try:
            cluster_col = hdr.index("集群")
        except ValueError:
            print(f"跳 {src.name}: R17 无'集群'列")
            continue
        total_col = next(
            (i for i, h in enumerate(hdr) if isinstance(h, str) and "合计件" in h),
            None,
        )
        if total_col is None:
            print(f"跳 {src.name}: R17 无'合计件'列")
            continue

        # 末尾追加 2 列
        approved_col_idx = len(hdr) + 1   # 1-based
        delta_col_idx = approved_col_idx + 1
        ws.cell(row=17, column=approved_col_idx, value=f"{offer_id} (审批件)")
        ws.cell(row=17, column=delta_col_idx, value="审批 vs solver (+/-)")
        for c in (
            ws.cell(row=17, column=approved_col_idx),
            ws.cell(row=17, column=delta_col_idx),
        ):
            c.font = bold
            c.alignment = Alignment(horizontal="center", wrap_text=True)

        # 已知 cluster 集合 (避免漏 4181 多出来的; 跳过末行 "合计" 小计)
        seen_clusters: set[str] = set()
        data_rows: list[int] = []
        for r in range(18, ws.max_row + 1):
            col1 = ws.cell(row=r, column=1).value
            cluster = ws.cell(row=r, column=cluster_col + 1).value
            if cluster is None:
                continue
            if isinstance(col1, str) and col1.strip() == "合计":
                continue  # 末行 "合计" 行
            seen_clusters.add(str(cluster))
            data_rows.append(r)
            solver_pieces = ws.cell(row=r, column=total_col + 1).value or 0
            approved_pieces = approved.get(str(cluster), 0)
            delta = approved_pieces - int(solver_pieces or 0)
            ws.cell(row=r, column=approved_col_idx, value=approved_pieces)
            ws.cell(row=r, column=delta_col_idx, value=delta)
            if delta > 0:
                ws.cell(row=r, column=delta_col_idx).fill = fill_green
            elif delta < 0:
                ws.cell(row=r, column=delta_col_idx).fill = fill_red

        # 4181 多出来的 cluster (allocation_xlsx 没的, 可能 hitl 加的)
        extras = [c for c in approved if c not in seen_clusters]
        if extras:
            for c in extras:
                row = ws.max_row + 1
                ws.cell(row=row, column=cluster_col + 1, value=c)
                ws.cell(row=row, column=approved_col_idx, value=approved[c])
                ws.cell(row=row, column=delta_col_idx, value=approved[c])
                ws.cell(row=row, column=delta_col_idx).fill = fill_yellow

        # 列宽
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(approved_col_idx)].width = 16
        ws.column_dimensions[get_column_letter(delta_col_idx)].width = 18

        # save
        out = src.with_name(src.stem + "_with_approved.xlsx")
        wb.save(out)
        solver_total = sum(
            int(ws.cell(row=r, column=total_col + 1).value or 0) for r in data_rows
        )
        approved_total = sum(approved.values())
        summaries.append({
            "src": src.name,
            "out": out.name,
            "offer_id": offer_id,
            "solver_total": solver_total,
            "approved_total": approved_total,
            "extra_clusters": extras,
        })
        print(f"✓ {src.name} → {out.name}: solver={solver_total} 审批={approved_total} (extras={extras or '无'})")

    print(f"\n总览:")
    for s in summaries:
        diff = s["approved_total"] - s["solver_total"]
        flag = "↑" if diff > 0 else ("↓" if diff < 0 else "=")
        print(f"  {s['offer_id']}: solver {s['solver_total']} → 审批 {s['approved_total']} {flag}{abs(diff)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
